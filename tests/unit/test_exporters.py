"""
Unit тесты для экспортеров данных.

Тестирование всех экспортеров: JSON, CSV, GraphML, HTML, PDF, Excel.
"""

import json
import pytest
from pathlib import Path
from xml.etree import ElementTree as ET

from webcrawler.export import (
    JSONExporter,
    CSVExporter,
    GraphMLExporter,
    HTMLExporter,
    PDFExporter,
    ExcelExporter,
    ExportFilter,
)


@pytest.mark.unit
@pytest.mark.asyncio
class TestJSONExporter:
    """Тесты для JSON экспортера."""

    async def test_json_export_basic(
        self, test_db_manager, sample_crawl_session, temp_test_dir
    ):
        """Тест базового JSON экспорта."""
        exporter = JSONExporter(test_db_manager)
        output_path = temp_test_dir / "export.json"

        await exporter.export(
            session_id=sample_crawl_session,
            output_path=str(output_path),
            pretty=True,
            compress=False
        )

        # Проверить файл создан
        assert output_path.exists()

        # Проверить валидность JSON
        with open(output_path, "r", encoding="utf-8") as f:
            data = json.load(f)

        # Проверить структуру
        assert "session" in data
        assert "urls" in data
        assert "links" in data
        assert "errors" in data
        assert "statistics" in data

        # Проверить данные сессии
        assert data["session"]["session_id"] == sample_crawl_session
        assert data["session"]["start_url"] == "https://example.com"

        # Проверить количество URL
        assert len(data["urls"]) == 5

    async def test_json_export_compact(
        self, test_db_manager, sample_crawl_session, temp_test_dir
    ):
        """Тест компактного JSON экспорта."""
        exporter = JSONExporter(test_db_manager)
        output_path = temp_test_dir / "export_compact.json"

        await exporter.export(
            session_id=sample_crawl_session,
            output_path=str(output_path),
            pretty=False,
            compress=False
        )

        # Компактный JSON должен быть меньше
        with open(output_path, "r") as f:
            content = f.read()
            # Не должно быть переводов строк (кроме одного в конце)
            assert content.count("\n") <= 1

    async def test_json_export_with_filter(
        self, test_db_manager, sample_crawl_session, temp_test_dir
    ):
        """Тест JSON экспорта с фильтром."""
        exporter = JSONExporter(test_db_manager)
        output_path = temp_test_dir / "export_filtered.json"

        # Фильтр только успешные запросы (200)
        exporter.add_filter(ExportFilter(status_codes=[200]))

        await exporter.export(
            session_id=sample_crawl_session,
            output_path=str(output_path),
            pretty=True
        )

        with open(output_path, "r", encoding="utf-8") as f:
            data = json.load(f)

        # Должно быть только URL со статусом 200
        for url in data["urls"]:
            assert url["status_code"] == 200

    async def test_json_export_nonexistent_session(
        self, test_db_manager, temp_test_dir
    ):
        """Тест экспорта несуществующей сессии."""
        exporter = JSONExporter(test_db_manager)
        output_path = temp_test_dir / "export_fail.json"

        with pytest.raises(ValueError, match="Session .* not found"):
            await exporter.export(
                session_id="nonexistent_session",
                output_path=str(output_path)
            )


@pytest.mark.unit
@pytest.mark.asyncio
class TestCSVExporter:
    """Тесты для CSV экспортера."""

    async def test_csv_export_basic(
        self, test_db_manager, sample_crawl_session, temp_test_dir
    ):
        """Тест базового CSV экспорта."""
        exporter = CSVExporter(test_db_manager)
        output_path = temp_test_dir / "export.csv"

        await exporter.export(
            session_id=sample_crawl_session,
            output_path=str(output_path),
            split_by_type=False
        )

        # Проверить файл создан
        assert output_path.exists()

        # Проверить содержимое
        with open(output_path, "r", encoding="utf-8") as f:
            lines = f.readlines()

        # Заголовок + 5 URL
        assert len(lines) >= 6

        # Проверить заголовок
        header = lines[0].strip()
        assert "url" in header.lower()
        assert "status_code" in header.lower()

    async def test_csv_export_split(
        self, test_db_manager, sample_crawl_session, temp_test_dir
    ):
        """Тест CSV экспорта с разделением на файлы."""
        exporter = CSVExporter(test_db_manager)
        output_path = temp_test_dir / "export.csv"

        await exporter.export(
            session_id=sample_crawl_session,
            output_path=str(output_path),
            split_by_type=True
        )

        # Проверить создание нескольких файлов
        assert (temp_test_dir / "export_urls.csv").exists()
        assert (temp_test_dir / "export_links.csv").exists()
        assert (temp_test_dir / "export_errors.csv").exists()


@pytest.mark.unit
@pytest.mark.asyncio
class TestGraphMLExporter:
    """Тесты для GraphML экспортера."""

    async def test_graphml_export_basic(
        self, test_db_manager, sample_crawl_session, temp_test_dir
    ):
        """Тест базового GraphML экспорта."""
        exporter = GraphMLExporter(test_db_manager)
        output_path = temp_test_dir / "graph.graphml"

        await exporter.export(
            session_id=sample_crawl_session,
            output_path=str(output_path),
            max_nodes=100
        )

        # Проверить файл создан
        assert output_path.exists()

        # Проверить валидность XML
        tree = ET.parse(output_path)
        root = tree.getroot()

        # Проверить namespace
        assert "graphml" in root.tag

        # Проверить наличие графа
        graphs = root.findall(".//{http://graphml.graphdrawing.org/xmlns}graph")
        assert len(graphs) > 0

    async def test_graphml_export_max_nodes(
        self, test_db_manager, sample_crawl_session, temp_test_dir
    ):
        """Тест GraphML экспорта с ограничением узлов."""
        exporter = GraphMLExporter(test_db_manager)
        output_path = temp_test_dir / "graph_limited.graphml"

        await exporter.export(
            session_id=sample_crawl_session,
            output_path=str(output_path),
            max_nodes=2  # Ограничить 2 узлами
        )

        tree = ET.parse(output_path)
        root = tree.getroot()

        # Подсчитать узлы
        nodes = root.findall(".//{http://graphml.graphdrawing.org/xmlns}node")
        assert len(nodes) <= 2


@pytest.mark.unit
@pytest.mark.asyncio
class TestHTMLExporter:
    """Тесты для HTML экспортера."""

    async def test_html_export_basic(
        self, test_db_manager, sample_crawl_session, temp_test_dir
    ):
        """Тест базового HTML экспорта."""
        exporter = HTMLExporter(test_db_manager)
        output_path = temp_test_dir / "report.html"

        await exporter.export(
            session_id=sample_crawl_session,
            output_path=str(output_path),
            theme="light",
            include_charts=True
        )

        # Проверить файл создан
        assert output_path.exists()

        # Проверить содержимое HTML
        with open(output_path, "r", encoding="utf-8") as f:
            html_content = f.read()

        # Проверить основные элементы
        assert "<!DOCTYPE html>" in html_content
        assert "<html" in html_content
        assert "Crawl Session Report" in html_content
        assert "Chart.js" in html_content  # CDN подключен
        assert sample_crawl_session in html_content

    async def test_html_export_dark_theme(
        self, test_db_manager, sample_crawl_session, temp_test_dir
    ):
        """Тест HTML экспорта с тёмной темой."""
        exporter = HTMLExporter(test_db_manager)
        output_path = temp_test_dir / "report_dark.html"

        await exporter.export(
            session_id=sample_crawl_session,
            output_path=str(output_path),
            theme="dark"
        )

        with open(output_path, "r", encoding="utf-8") as f:
            html_content = f.read()

        # Проверить тёмные цвета
        assert "#1a1a2e" in html_content or "#2d2d44" in html_content


@pytest.mark.unit
@pytest.mark.asyncio
class TestPDFExporter:
    """Тесты для PDF экспортера."""

    async def test_pdf_exporter_available(self, test_db_manager):
        """Тест доступности PDF экспортера."""
        exporter = PDFExporter(test_db_manager)

        # Может быть недоступен если reportlab не установлен
        # Это нормально, экспортер должен корректно обрабатывать это
        assert hasattr(exporter, "reportlab_available")

    @pytest.mark.skipif(
        not __import__("importlib").util.find_spec("reportlab"),
        reason="reportlab not installed"
    )
    async def test_pdf_export_basic(
        self, test_db_manager, sample_crawl_session, temp_test_dir
    ):
        """Тест базового PDF экспорта."""
        exporter = PDFExporter(test_db_manager)
        output_path = temp_test_dir / "report.pdf"

        await exporter.export(
            session_id=sample_crawl_session,
            output_path=str(output_path),
            page_size="A4",
            orientation="portrait"
        )

        # Проверить файл создан
        assert output_path.exists()

        # Проверить что это PDF (magic bytes)
        with open(output_path, "rb") as f:
            header = f.read(4)
            assert header == b"%PDF"


@pytest.mark.unit
@pytest.mark.asyncio
class TestExcelExporter:
    """Тесты для Excel экспортера."""

    async def test_excel_exporter_available(self, test_db_manager):
        """Тест доступности Excel экспортера."""
        exporter = ExcelExporter(test_db_manager)

        # Может быть недоступен если openpyxl не установлен
        assert hasattr(exporter, "openpyxl_available")

    @pytest.mark.skipif(
        not __import__("importlib").util.find_spec("openpyxl"),
        reason="openpyxl not installed"
    )
    async def test_excel_export_basic(
        self, test_db_manager, sample_crawl_session, temp_test_dir
    ):
        """Тест базового Excel экспорта."""
        exporter = ExcelExporter(test_db_manager)
        output_path = temp_test_dir / "report.xlsx"

        await exporter.export(
            session_id=sample_crawl_session,
            output_path=str(output_path),
            include_charts=True,
            include_links=True,
            include_errors=True
        )

        # Проверить файл создан
        assert output_path.exists()

        # Проверить структуру Excel
        from openpyxl import load_workbook

        wb = load_workbook(output_path)

        # Проверить листы
        assert "Summary" in wb.sheetnames
        assert "URLs" in wb.sheetnames
        assert "Links" in wb.sheetnames
        assert "Errors" in wb.sheetnames
        assert "Status Codes" in wb.sheetnames
        assert "Content Types" in wb.sheetnames


@pytest.mark.unit
@pytest.mark.asyncio
class TestExportFilter:
    """Тесты для фильтров экспорта."""

    async def test_filter_by_status_codes(
        self, test_db_manager, sample_crawl_session
    ):
        """Тест фильтрации по статус-кодам."""
        exporter = JSONExporter(test_db_manager)

        # Фильтр только 200 и 404
        exporter.add_filter(ExportFilter(status_codes=[200, 404]))

        # Получить данные
        data = await exporter._fetch_session_data(sample_crawl_session)

        # Проверить что все URL имеют статус 200 или 404
        for url in data["urls"]:
            assert url.status_code in [200, 404]

    async def test_filter_by_depth(
        self, test_db_manager, sample_crawl_session
    ):
        """Тест фильтрации по глубине."""
        exporter = JSONExporter(test_db_manager)

        # Только глубина 0
        exporter.add_filter(ExportFilter(max_depth=0))

        data = await exporter._fetch_session_data(sample_crawl_session)

        # Проверить что все URL имеют глубину 0
        for url in data["urls"]:
            assert url.depth == 0

    async def test_filter_by_content_type(
        self, test_db_manager, sample_crawl_session
    ):
        """Тест фильтрации по типу контента."""
        exporter = JSONExporter(test_db_manager)

        # Только HTML
        exporter.add_filter(ExportFilter(content_types=["text/html"]))

        data = await exporter._fetch_session_data(sample_crawl_session)

        # Проверить что все URL имеют content_type text/html
        for url in data["urls"]:
            assert url.content_type == "text/html"

    async def test_multiple_filters(
        self, test_db_manager, sample_crawl_session
    ):
        """Тест комбинации фильтров."""
        exporter = JSONExporter(test_db_manager)

        # Статус 200 И глубина 1 И HTML
        exporter.add_filter(
            ExportFilter(
                status_codes=[200],
                max_depth=1,
                content_types=["text/html"]
            )
        )

        data = await exporter._fetch_session_data(sample_crawl_session)

        # Проверить что все условия выполнены
        for url in data["urls"]:
            assert url.status_code == 200
            assert url.depth <= 1
            assert url.content_type == "text/html"
